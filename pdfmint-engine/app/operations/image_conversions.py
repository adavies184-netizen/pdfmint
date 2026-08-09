from __future__ import annotations

import base64
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytesseract
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageSequence
from pillow_heif import register_heif_opener

register_heif_opener()


ROUTES = {
    "image-to-jpg": ("image", "jpg"), "image-to-png": ("image", "png"),
    "image-to-svg": ("image", "svg"), "image-to-gif": ("image", "gif"),
    "image-to-word": ("image", "docx"), "image-to-excel": ("image", "xlsx"),
    "heic-to-jpg": ("heic", "jpg"), "heic-to-png": ("heic", "png"),
    "png-to-jpg": ("png", "jpg"), "jpg-to-png": ("jpg", "png"),
    "png-to-eps": ("png", "eps"), "png-to-ico": ("png", "ico"),
    "jpeg-to-eps": ("jpeg", "eps"), "webp-to-jpg": ("webp", "jpg"),
    "jpeg-to-png": ("jpeg", "png"), "svg-to-png": ("svg", "png"),
    "svg-to-dxf": ("svg", "dxf"), "eps-to-svg": ("eps", "svg"),
    "jfif-to-jpg": ("jfif", "jpg"), "avif-to-jpg": ("avif", "jpg"),
    "docx-to-jpg": ("docx", "jpg"), "doc-to-jpg": ("doc", "jpg"),
    "word-to-jpg": ("word", "jpg"), "html-to-jpg": ("html", "jpg"),
}

SOURCE_SUFFIXES = {
    "image": {".avif", ".bmp", ".gif", ".heic", ".heif", ".jfif", ".jpg", ".jpeg", ".png", ".svg", ".tif", ".tiff", ".webp"},
    "heic": {".heic", ".heif"}, "png": {".png"}, "jpg": {".jpg", ".jpeg"},
    "jpeg": {".jpg", ".jpeg"}, "webp": {".webp"}, "svg": {".svg"},
    "eps": {".eps"}, "jfif": {".jfif", ".jpg", ".jpeg"}, "avif": {".avif"},
    "docx": {".docx"}, "doc": {".doc"}, "word": {".doc", ".docx"},
    "html": {".html", ".htm"},
}


def _run(args: list[str], cwd: Path | None = None) -> None:
    subprocess.run(args, cwd=cwd, check=True, capture_output=True, text=True, timeout=180)


def _flatten(image: Image.Image) -> Image.Image:
    image = image.convert("RGBA")
    background = Image.new("RGBA", image.size, "white")
    background.alpha_composite(image)
    return background.convert("RGB")


def _readable_image(input_path: Path) -> Path:
    if input_path.suffix.lower() != ".svg":
        return input_path
    raster = input_path.with_name("svg-source.png")
    _run(["rsvg-convert", "-o", str(raster), str(input_path)])
    return raster


def _raster(input_path: Path, output: Path, target: str) -> None:
    with Image.open(_readable_image(input_path)) as opened:
        image = ImageSequence.Iterator(opened).__next__()
        if target in {"jpg", "jpeg", "eps"}:
            image = _flatten(image)
        elif target == "png":
            image = image.convert("RGBA")
        elif target == "ico":
            image = image.convert("RGBA")
            image.thumbnail((256, 256), Image.Resampling.LANCZOS)
        image.save(output, format={"jpg": "JPEG", "png": "PNG", "gif": "GIF", "eps": "EPS", "ico": "ICO"}[target], quality=92)


def _image_to_svg(input_path: Path, output: Path) -> None:
    if input_path.suffix.lower() == ".svg":
        shutil.copy2(input_path, output)
        return
    mime = Image.MIME.get(Image.open(input_path).format, "image/png")
    data = base64.b64encode(input_path.read_bytes()).decode("ascii")
    with Image.open(input_path) as image:
        width, height = image.size
    output.write_text(
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
        f'<image width="{width}" height="{height}" href="data:{mime};base64,{data}"/></svg>',
        encoding="utf-8",
    )


def _ocr_docx(input_path: Path, output: Path) -> None:
    text = pytesseract.image_to_string(Image.open(_readable_image(input_path)))
    document = Document()
    for block in text.split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())
    document.save(output)


def _ocr_xlsx(input_path: Path, output: Path) -> None:
    data = pytesseract.image_to_data(Image.open(_readable_image(input_path)), output_type=pytesseract.Output.DICT)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extracted text"
    rows: dict[tuple[int, int, int], list[str]] = {}
    for index, word in enumerate(data["text"]):
        word = word.strip()
        if word:
            key = (data["page_num"][index], data["block_num"][index], data["line_num"][index])
            rows.setdefault(key, []).append(word)
    for words in rows.values():
        sheet.append(words)
    workbook.save(output)


def _office_to_jpg(input_path: Path, workspace: Path, output: Path) -> Path:
    _run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(workspace), str(input_path)])
    pdf = next(workspace.glob("*.pdf"))
    prefix = workspace / "page"
    _run(["pdftoppm", "-jpeg", "-r", "150", str(pdf), str(prefix)])
    pages = sorted(workspace.glob("page-*.jpg"))
    if len(pages) == 1:
        shutil.move(pages[0], output)
        return output
    archive = output.with_suffix(".zip")
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for index, page in enumerate(pages, 1):
            bundle.write(page, f"page-{index}.jpg")
    return archive


def convert_image_route(operation: str, input_path: Path, workspace: Path, base_name: str) -> tuple[Path, str]:
    source, target = ROUTES[operation]
    if input_path.suffix.lower() not in SOURCE_SUFFIXES[source]:
        raise ValueError(f"This converter does not accept {input_path.suffix or 'that'} files.")
    output = workspace / f"{base_name}.{target}"
    if operation in {"image-to-word"}:
        _ocr_docx(input_path, output)
    elif operation in {"image-to-excel"}:
        _ocr_xlsx(input_path, output)
    elif operation in {"image-to-svg"}:
        _image_to_svg(input_path, output)
    elif operation == "svg-to-png":
        _run(["rsvg-convert", "-o", str(output), str(input_path)])
    elif operation == "svg-to-dxf":
        _run(["soffice", "--headless", "--convert-to", "dxf", "--outdir", str(workspace), str(input_path)])
        converted = next((p for p in workspace.glob("*.dxf") if p != output), output)
        if converted != output:
            shutil.move(converted, output)
    elif operation == "eps-to-svg":
        intermediate = workspace / "source.pdf"
        _run(["gs", "-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", f"-sOutputFile={intermediate}", str(input_path)])
        _run(["pdftocairo", "-svg", str(intermediate), str(output)])
    elif operation in {"docx-to-jpg", "doc-to-jpg", "word-to-jpg", "html-to-jpg"}:
        output = _office_to_jpg(input_path, workspace, output)
    else:
        _raster(input_path, output, target)
    if not output.exists() or output.stat().st_size == 0:
        raise RuntimeError("The converted file was not produced.")
    return output, output.suffix.lower()
